"""A claim answered line by line, and every refused line followed.

Asked as step 5 of the hospital insurance plan — the claims desk: a batch by
payer and cycle day, the lines from every department with their approvals,
the payer's refusals per line with the reason, resubmission, and what the
payer actually paid against what was claimed. What is held here:

* a claim carries its lines — each covered bill line, with its department
  and its approval number;
* the payer's answer is written per line; anything refused says why, and
  the claim's accepted figure is the lines' sum;
* a refused remainder is sent again on a resubmission claim that points
  back at it, or written off by somebody who says so — never both;
* the claim reads claimed, accepted, paid and still open; a payment short
  of what was accepted is said;
* each payer's next batch runs from the day after its last claimed period
  to its contract's cycle day;
* the claim downloads as a sheet, one row per line.
"""
import os
import sys
from datetime import date, timedelta
from io import BytesIO

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import pytest  # noqa: E402

from app.utils.clock import local_today  # noqa: E402
from tests.test_a_contract_that_knows_the_departments import (  # noqa: E402,F401
    _rule, insured)


def _bill(c, lines, number="INV-1"):
    from app.models import Invoice, InvoiceItem, Patient, Service
    from app.utils import billing

    with c["app"].app_context():
        db = c["db"]
        inv = Invoice(patient_id=c["ids"]["child"], invoice_number=number,
                      invoice_date=local_today(), status="unpaid")
        db.session.add(inv)
        db.session.flush()
        for key, price in lines:
            svc = db.session.get(Service, c["ids"][key])
            db.session.add(InvoiceItem(invoice_id=inv.id, service_id=svc.id,
                                       description=svc.name, quantity=1, unit_price=price))
        db.session.flush()
        db.session.refresh(inv)
        billing.apply_coverage(inv, db.session.get(Patient, c["ids"]["child"]))
        db.session.commit()
        return inv.id


def _claim(c):
    from app.models import Claim

    day = local_today().isoformat()
    c["sign_in"]("boss").post("/finance/claims/create", data={
        "payer_id": c["ids"]["payer"], "date_from": day, "date_to": day})
    with c["app"].app_context():
        return Claim.query.order_by(Claim.id.desc()).first().id


def _submit(c, claim_id):
    c["sign_in"]("boss").post(f"/finance/claim/{claim_id}/action", data={"action": "submit"})


def _lines(c, claim_id):
    from app.models import Claim

    with c["app"].app_context():
        claim = c["db"].session.get(Claim, claim_id)
        return [(ln.id, ln.description, ln.amount, ln.setting, ln.decision, ln.accepted)
                for ln in claim.lines]


@pytest.fixture()
def claimed(insured):
    _rule(insured, setting="any", scope="all", coverage_type="percent", coverage_value=100)
    _bill(insured, [("exam", 200), ("cbc", 100)])
    claim = _claim(insured)
    insured["ids"]["claim"] = claim
    return insured


def test_a_claim_carries_its_lines_with_their_department(claimed):
    lines = _lines(claimed, claimed["ids"]["claim"])
    assert [(d, a, s) for _i, d, a, s, _x, _y in lines] == [
        ("كشف", 200.0, "outpatient"), ("صورة دم", 100.0, "outpatient")]
    page = claimed["sign_in"]("boss").get(f"/finance/claim/{claimed['ids']['claim']}").get_data(as_text=True)
    assert "data-claim-lines" in page and "data-claim-sums" in page


def test_the_payer_answers_by_the_line_and_a_refusal_says_why(claimed):
    from app.models import Claim

    claim_id = claimed["ids"]["claim"]
    _submit(claimed, claim_id)
    (exam, *_), (cbc, *_) = _lines(claimed, claim_id)
    boss = claimed["sign_in"]("boss")
    # A refusal with no reason is refused.
    boss.post(f"/finance/claim/{claim_id}/decide", data={f"accepted_{cbc}": "0"})
    with claimed["app"].app_context():
        assert claimed["db"].session.get(Claim, claim_id).status == "submitted"
    boss.post(f"/finance/claim/{claim_id}/decide",
              data={f"accepted_{cbc}": "0", f"reason_{cbc}": "مش في البوليصة"})
    with claimed["app"].app_context():
        claim = claimed["db"].session.get(Claim, claim_id)
        assert claim.status == "approved" and claim.approved_amount == 200.0
        assert [(ln.decision, ln.refusal_reason) for ln in claim.lines] == [
            ("accepted", None), ("refused", "مش في البوليصة")]


def test_a_refused_line_is_sent_again_on_a_claim_that_points_back(claimed):
    from app.models import Claim, ClaimLine

    claim_id = claimed["ids"]["claim"]
    _submit(claimed, claim_id)
    (exam, *_), (cbc, *_) = _lines(claimed, claim_id)
    boss = claimed["sign_in"]("boss")
    boss.post(f"/finance/claim/{claim_id}/decide",
              data={f"accepted_{exam}": "150", f"reason_{exam}": "تعريفة",
                    f"accepted_{cbc}": "0", f"reason_{cbc}": "محتاج موافقة"})
    answer = boss.post(f"/finance/claim/{claim_id}/refused",
                       data={"line_id": [str(exam), str(cbc)], "action": "resubmit"})
    with claimed["app"].app_context():
        db = claimed["db"]
        again = Claim.query.filter_by(resubmission_of_id=claim_id).one()
        assert f"/finance/claim/{again.id}" in answer.headers["Location"]
        assert again.total_amount == 150.0 and again.status == "draft"
        assert sorted(ln.amount for ln in again.lines) == [50.0, 100.0]
        assert {ln.resubmit_of_id for ln in again.lines} == {exam, cbc}
        assert all(db.session.get(ClaimLine, i).resubmitted_in_id == again.id for i in (exam, cbc))
    # Already sent again: it cannot also be written off.
    boss.post(f"/finance/claim/{claim_id}/refused", data={"line_id": [str(cbc)], "action": "write_off"})
    with claimed["app"].app_context():
        assert not claimed["db"].session.get(ClaimLine, cbc).written_off


def test_a_refused_line_written_off_and_a_short_payment_said(claimed):
    from app.models import ClaimLine

    claim_id = claimed["ids"]["claim"]
    _submit(claimed, claim_id)
    (exam, *_), (cbc, *_) = _lines(claimed, claim_id)
    boss = claimed["sign_in"]("boss")
    boss.post(f"/finance/claim/{claim_id}/decide",
              data={f"accepted_{cbc}": "0", f"reason_{cbc}": "مستثنى"})
    boss.post(f"/finance/claim/{claim_id}/refused", data={"line_id": [str(cbc)], "action": "write_off"})
    with claimed["app"].app_context():
        assert claimed["db"].session.get(ClaimLine, cbc).written_off is True
    boss.post(f"/finance/claim/{claim_id}/action",
              data={"action": "pay", "paid_amount": "180", "payment_method": "transfer"})
    page = boss.get(f"/finance/claim/{claim_id}").get_data(as_text=True)
    assert "data-paid-short" in page
    from app.models import Claim
    from app.utils import claims_desk

    with claimed["app"].app_context():
        sums = claims_desk.totals(claimed["db"].session.get(Claim, claim_id))
        assert (sums["claimed"], sums["accepted"], sums["paid"], sums["open"],
                sums["written_off"], sums["short"]) == (300.0, 200.0, 180.0, 0.0, 100.0, 20.0)


def test_the_next_batch_runs_to_the_cycle_day(insured):
    from app.models import Claim, PayerContract, PayerEntity
    from app.utils import claims_desk

    with insured["app"].app_context():
        db = insured["db"]
        contract = db.session.get(PayerContract, insured["ids"]["contract"])
        contract.cycle_day = 25
        payer = db.session.get(PayerEntity, insured["ids"]["payer"])
        db.session.add(Claim(claim_number="CLM-X-1", payer_id=payer.id,
                             date_from=date(2026, 8, 1), date_to=date(2026, 8, 25),
                             status="paid"))
        db.session.commit()
        assert claims_desk.next_batch(payer, today=date(2026, 9, 30)) == (
            date(2026, 8, 26), date(2026, 9, 25))
        # Before the month's cycle day, last month's cutoff — already claimed.
        assert claims_desk.next_batch(payer, today=date(2026, 9, 10)) is None
    page = insured["sign_in"]("boss").get("/finance/claims").get_data(as_text=True)
    assert "data-claims-board" in page


def test_the_claim_downloads_as_a_sheet_one_row_per_line(claimed):
    from openpyxl import load_workbook

    raw = claimed["sign_in"]("boss").get(f"/finance/claim/{claimed['ids']['claim']}/sheet")
    assert raw.status_code == 200
    ws = load_workbook(BytesIO(raw.data)).worksheets[0]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert [(r[5], r[7]) for r in rows] == [("كشف", 200), ("صورة دم", 100)]


def test_a_resubmission_taken_back_opens_the_refusal_again(claimed):
    from app.models import Claim, ClaimLine

    claim_id = claimed["ids"]["claim"]
    _submit(claimed, claim_id)
    (exam, *_), (cbc, *_) = _lines(claimed, claim_id)
    boss = claimed["sign_in"]("boss")
    boss.post(f"/finance/claim/{claim_id}/decide",
              data={f"accepted_{cbc}": "0", f"reason_{cbc}": "محتاج موافقة"})
    boss.post(f"/finance/claim/{claim_id}/refused", data={"line_id": [str(cbc)], "action": "resubmit"})
    with claimed["app"].app_context():
        again = Claim.query.filter_by(resubmission_of_id=claim_id).one().id
    boss.post(f"/finance/claim/{again}/action", data={"action": "delete"})
    with claimed["app"].app_context():
        line = claimed["db"].session.get(ClaimLine, cbc)
        assert line.resubmitted_in_id is None and line.open_refusal
