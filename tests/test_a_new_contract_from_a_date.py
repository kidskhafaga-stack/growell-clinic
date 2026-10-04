"""A new contract from a date, its prices raised or uploaded, and a reminder
before the old one ends.

Asked as *«لو الادارة عايز تعمل عقد جديد مثلاً من تاريخ كذا الى كذا … انا كنت
عامل العقد من كذا لكذا علشان بعد كده يتنفذ لوحده»*, *«هزود 20% على الكشف
هزود 30% على كذا … او انزل شيت للعقد … واعدل على اسعاره وبعد كده ارفعه»* and
*«عايز نظام للتذكير بانتهاء عقد»*. What is held here:

* a renewal with new dates can raise every price in the same step, and
  bills by itself from its first day;
* a raise by service beats one by category beats «everything»; rounding is
  the hospital's;
* a raise and a sheet both stop at a preview — nothing is written until it
  is confirmed, and the confirm writes exactly what was shown;
* the sheet round-trips: a price changed, a service newly covered, a
  service out of the contract, a row nobody recognises;
* the reminder: days left, renewal ready, and a contract that ended with no
  renewal while members still carry cards.
"""
import os
import sys
from datetime import timedelta
from io import BytesIO

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import pytest  # noqa: E402

from app.utils.clock import local_today  # noqa: E402


@pytest.fixture()
def priced(clinic):
    from app.models import (PatientCoverage, PayerContract, PayerContractRate,
                            PayerEntity, Service)

    db = clinic["db"]
    with clinic["app"].app_context():
        today = local_today()
        payer = PayerEntity(name="شركة تأمين", entity_type="insurance", is_active=True)
        db.session.add(payer)
        db.session.flush()
        old = PayerContract(payer_id=payer.id, number="C-1", is_active=True,
                            start_date=today - timedelta(days=300),
                            end_date=today + timedelta(days=10))
        db.session.add(old)
        cbc = Service(name="صورة دم", code="LAB1", category="lab", price=100, is_active=True)
        film = Service(name="أشعة صدر", code="RAD1", category="radiology", price=300, is_active=True)
        db.session.add_all([cbc, film])
        db.session.flush()
        db.session.add(PayerContractRate(contract_id=old.id, service_id=clinic["ids"]["exam"],
                                         special_price=200, coverage_type="percent",
                                         coverage_value=80))
        db.session.add(PayerContractRate(contract_id=old.id, service_id=cbc.id,
                                         coverage_type="percent", coverage_value=100))
        db.session.add(PatientCoverage(patient_id=clinic["ids"]["child"], payer_id=payer.id,
                                       membership_number="M-1", is_active=True))
        db.session.commit()
        clinic["ids"].update(payer=payer.id, old=old.id, cbc=cbc.id, film=film.id)
    return clinic


def _rates(c, contract_id):
    from app.models import PayerContract

    with c["app"].app_context():
        row = c["db"].session.get(PayerContract, contract_id)
        return {r.service_id: (r.special_price, r.coverage_type if r.coverage_value else "none",
                               r.coverage_value or 0) for r in row.rates}


def _renew(c, **extra):
    from app.models import PayerContract

    today = local_today()
    c["sign_in"]("boss").post(
        f"/finance/contract/{c['ids']['old']}/copy",
        data=dict(start_date=(today + timedelta(days=11)).isoformat(),
                  end_date=(today + timedelta(days=376)).isoformat(), **extra))
    with c["app"].app_context():
        return PayerContract.query.order_by(PayerContract.id.desc()).first().id


# -------------------------------------------------------------- renewal --
def test_a_renewal_raises_every_price_in_the_same_step(priced):
    new = _renew(priced, raise_percent="20")
    got = _rates(priced, new)
    assert got[priced["ids"]["exam"]] == (240.0, "percent", 80.0)
    assert got[priced["ids"]["cbc"]] == (120.0, "percent", 100.0), "the clinic price was not raised"
    assert priced["ids"]["film"] not in got, "a raise put a new service on the contract"


def test_the_renewal_bills_by_itself_from_its_first_day(priced):
    from app.models import PayerEntity

    new = _renew(priced)
    with priced["app"].app_context():
        payer = priced["db"].session.get(PayerEntity, priced["ids"]["payer"])
        today = local_today()
        assert payer.active_contract(today).id == priced["ids"]["old"]
        assert payer.active_contract(today + timedelta(days=11)).id == new


# ----------------------------------------------------------- the raise --
def test_the_most_specific_raise_decides_and_nothing_is_saved_before_confirm(priced):
    boss = priced["sign_in"]("boss")
    before = _rates(priced, priced["ids"]["old"])
    page = boss.post(f"/finance/contract/{priced['ids']['old']}/adjust", data={
        "target": ["all", "cat:lab", f"svc:{priced['ids']['exam']}"],
        "kind": ["percent", "percent", "amount"],
        "value": ["10", "30", "50"], "step": "5"}).get_data(as_text=True)
    assert 'data-preview-row="changed"' in page and "data-preview-confirm" in page
    assert _rates(priced, priced["ids"]["old"]) == before, "the preview wrote"

    # The confirm posts the very list that was shown.
    from app.utils import contract_renewal as renewal
    from app.models import PayerContract, Service

    with priced["app"].app_context():
        c = priced["db"].session.get(PayerContract, priced["ids"]["old"])
        services = Service.query.filter_by(is_active=True).order_by(Service.name).all()
        proposal = renewal.raised(c, services, [("all", "percent", 10), ("cat:lab", "percent", 30),
                                                (f"svc:{priced['ids']['exam']}", "amount", 50)], 5)
    form = {}
    for sid, (price, ctype, value) in proposal.items():
        form[f"price_{sid}"] = "" if price is None else str(price)
        form[f"type_{sid}"] = ctype
        form[f"value_{sid}"] = "" if ctype == "none" else str(value)
    boss.post(f"/finance/contract/{priced['ids']['old']}/rates", data=form)
    got = _rates(priced, priced["ids"]["old"])
    assert got[priced["ids"]["exam"]] == (250.0, "percent", 80.0)
    assert got[priced["ids"]["cbc"]] == (130.0, "percent", 100.0)


def test_a_raise_with_no_value_is_refused(priced):
    answer = priced["sign_in"]("boss").post(f"/finance/contract/{priced['ids']['old']}/adjust",
                                            data={"target": ["all"], "kind": ["percent"],
                                                  "value": [""]})
    assert answer.status_code == 302


# ----------------------------------------------------------- the sheet --
def _sheet(priced):
    from openpyxl import load_workbook

    raw = priced["sign_in"]("boss").get(f"/finance/contract/{priced['ids']['old']}/sheet")
    assert raw.status_code == 200
    return load_workbook(BytesIO(raw.data))


def test_the_sheet_round_trips_with_every_kind_of_change(priced):
    book = _sheet(priced)
    ws = book.worksheets[0]
    rows = {r[0].value: r for r in ws.iter_rows(min_row=2)}
    exam_code = None
    for r in ws.iter_rows(min_row=2):
        if r[1].value == "كشف":
            r[4].value = 260                      # a price changed
            exam_code = r
    rows["LAB1"][4].value = None                  # out of the contract
    rows["LAB1"][5].value = "غير مغطى"
    rows["RAD1"][4].value = 320                   # newly covered
    rows["RAD1"][5].value = "نسبة"
    rows["RAD1"][6].value = 70
    ws.append(["X-404", "خدمة مش موجودة", "", 0, 50, "نسبة", 50])
    assert exam_code is not None
    buf = BytesIO()
    book.save(buf)
    buf.seek(0)
    boss = priced["sign_in"]("boss")
    page = boss.post(f"/finance/contract/{priced['ids']['old']}/sheet",
                     data={"sheet": (buf, "c.xlsx")},
                     content_type="multipart/form-data").get_data(as_text=True)
    for kind in ("changed", "added", "removed"):
        assert f'data-preview-row="{kind}"' in page, kind
    assert "data-preview-skipped" in page
    assert _rates(priced, priced["ids"]["old"])[priced["ids"]["exam"]][0] == 200.0


def test_a_file_that_is_not_a_sheet_is_refused(priced):
    answer = priced["sign_in"]("boss").post(
        f"/finance/contract/{priced['ids']['old']}/sheet",
        data={"sheet": (BytesIO(b"not a workbook"), "x.xlsx")},
        content_type="multipart/form-data")
    assert answer.status_code == 302


# --------------------------------------------------------- the reminder --
def test_the_reminder_says_days_left_and_whether_a_renewal_is_ready(priced):
    from app.models import PayerContract
    from app.utils import contract_renewal as renewal

    with priced["app"].app_context():
        old = priced["db"].session.get(PayerContract, priced["ids"]["old"])
        found = renewal.state(old)
        assert (found["key"], found["days"]) == ("ending", 10)
        ending, lapsed = renewal.needing_attention()
        assert [c.id for c in ending] == [old.id] and lapsed == []
    page = priced["sign_in"]("boss").get("/finance/payers").get_data(as_text=True)
    assert 'data-contract-due="ending"' in page
    _renew(priced)
    with priced["app"].app_context():
        old = priced["db"].session.get(PayerContract, priced["ids"]["old"])
        assert renewal.state(old)["key"] == "renewed"
        assert renewal.needing_attention() == ([], [])


def test_a_contract_that_lapsed_with_members_is_said_louder(priced):
    from app.models import PayerContract
    from app.utils import contract_renewal as renewal

    with priced["app"].app_context():
        old = priced["db"].session.get(PayerContract, priced["ids"]["old"])
        old.end_date = local_today() - timedelta(days=3)
        priced["db"].session.commit()
        found = renewal.state(old)
        assert found["key"] == "lapsed" and found["members"] == 1
        assert [c.id for c in renewal.needing_attention()[1]] == [old.id]


def test_the_reminder_window_is_the_clinics(priced):
    from app.models import PayerContract, Setting
    from app.utils import contract_renewal as renewal

    with priced["app"].app_context():
        Setting.set("contract_remind_days", "5")
        priced["db"].session.commit()
        assert renewal.state(priced["db"].session.get(PayerContract, priced["ids"]["old"])) is None
