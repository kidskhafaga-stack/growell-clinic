"""A test defined in full — from a sheet, or one step at a time.

Asked as *«التفاصيل دي متاح رفعها او انزال نموذج تضاف بشكل كامل وترفع او
اضافة واحد لواحد بالخطوات المطلوبة المنطقية للبرنامج علشان التحليل يتسعّر
سعر وتكلفة وكل حاجه»*. What is held here:

* an **empty template** with every column the import reads, each header
  noting what goes in it — and no sample row, which would be read back in;
* the sheet's **price** gives a test with no price one of its own, and never
  changes a price a test already has; its **cost** fills an empty cost; its
  **consumables** are matched to the store, and a name the store does not
  know is counted, never created;
* the catalogue exported comes back in with the same price and cost;
* a test added on the screen lands on its own page, with the steps it still
  needs listed in order, and the sample/tube/time step saves.
"""
import io
import os
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import pytest  # noqa: E402


@pytest.fixture()
def lab(clinic):
    from app.models import Setting, StoreItem

    with clinic["app"].app_context():
        db = clinic["db"]
        Setting.set("mod_enabled:labs", "1")
        strip = StoreItem(name="شريط سكر", item_code="STR-01", is_active=True)
        db.session.add(strip)
        db.session.commit()
        clinic["ids"]["strip"] = strip.id
    return clinic


def _sheet(rows, header=None):
    import openpyxl

    from app.utils.lab_import import EXPORT_COLUMNS

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(list(header or EXPORT_COLUMNS))
    for row in rows:
        sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    out.seek(0)
    return out


def _row(**cells):
    from app.utils.lab_import import EXPORT_COLUMNS

    return [cells.get(c, "") for c in EXPORT_COLUMNS]


def _import(lab, stream):
    from app.utils import lab_import

    with lab["app"].app_context():
        plan = lab_import.read(stream)
        lab_import.match(plan)
        counts = lab_import.apply(plan, lab_import.links_from(plan, {}))
        lab["db"].session.commit()
        return counts


def test_the_empty_template_has_every_column_and_no_sample_row(lab):
    import openpyxl

    from app.utils.lab_import import EXPORT_COLUMNS, HEADERS, _header_map

    answer = lab["sign_in"]("boss").get("/labs/tests/template")
    assert answer.status_code == 200
    book = openpyxl.load_workbook(io.BytesIO(answer.data))
    sheet = book.active
    header = [c.value for c in sheet[1]]
    assert header == list(EXPORT_COLUMNS)
    assert sheet.max_row == 1, "a sample row would be read back in as a test"
    assert sheet["A1"].comment is not None
    read = _header_map(header)
    for column in ("price", "cost", "consumable", "consumable_qty"):
        assert column in read and column in HEADERS


def test_the_sheets_price_cost_and_consumables_fill_a_new_test(lab):
    from app.models import Investigation

    counts = _import(lab, _sheet([
        _row(**{"Test Name": "Random Glucose", "Arabic Test Name": "سكر عشوائي",
                "Price": 60, "Cost": 12.5, "Consumable": "STR-01", "Consumable Qty": 1}),
        _row(**{"Test Name": "Random Glucose", "Consumable": "لانسيت مش موجود",
                "Consumable Qty": 2}),
    ]))
    assert counts["priced"] == 1 and counts["unmatched"] == 1
    with lab["app"].app_context():
        row = Investigation.query.filter_by(name_en="Random Glucose").one()
        assert row.service is not None and row.service.price == 60
        assert row.service.category == "lab"
        assert row.cost == 12.5
        assert [(c.store_item_id, c.quantity) for c in row.lab_consumables] == [
            (lab["ids"]["strip"], 1)]


def test_the_sheet_never_changes_a_price_a_test_already_has(lab):
    from app.models import Investigation

    with lab["app"].app_context():
        db = lab["db"]
        db.session.add(Investigation(name_ar="سكر عشوائي", name_en="Random Glucose",
                                     kind="lab", service_id=lab["ids"]["exam"],
                                     cost=9, is_active=True))
        db.session.commit()
    counts = _import(lab, _sheet([_row(**{"Test Name": "Random Glucose",
                                          "Price": 999, "Cost": 50})]))
    assert counts["priced"] == 0
    with lab["app"].app_context():
        row = Investigation.query.filter_by(name_en="Random Glucose").one()
        assert row.service_id == lab["ids"]["exam"] and row.service.price == 200
        assert row.cost == 9


def test_what_is_exported_comes_back_with_its_price_and_cost(lab):
    from app.models import Investigation, LabConsumable
    from app.utils import lab_import

    with lab["app"].app_context():
        db = lab["db"]
        test = Investigation(name_ar="سكر", name_en="Glucose", kind="lab",
                             service_id=lab["ids"]["exam"], cost=7, is_active=True)
        db.session.add(test)
        db.session.flush()
        db.session.add(LabConsumable(investigation_id=test.id,
                                     store_item_id=lab["ids"]["strip"], quantity=3))
        db.session.commit()
        data = lab_import.export([test])
        plan = lab_import.read(io.BytesIO(data))
    found = plan["tests"][0]
    assert found["price"] == 200 and found["cost"] == 7
    assert found["consumables"] == [["STR-01", 3]]


def test_a_test_added_on_the_screen_lands_on_its_steps(lab):
    from app.models import Investigation

    boss = lab["sign_in"]("boss")
    answer = boss.post("/labs/tests/add", data={"name_ar": "وظائف كلى", "kind": "lab",
                                                "in_house": "1"})
    with lab["app"].app_context():
        row = Investigation.query.filter_by(name_ar="وظائف كلى").one()
        test_id = row.id
    assert answer.headers["Location"].endswith(f"/labs/tests/{test_id}/ranges")
    page = boss.get(f"/labs/tests/{test_id}/ranges").get_data(as_text=True)
    assert 'data-step="name" data-done="1"' in page
    assert 'data-step="sample" data-done="0"' in page
    assert 'data-step="price" data-done="0"' in page

    boss.post(f"/labs/tests/{test_id}/details",
              data={"sample_type": "دم", "tube": "أحمر", "tat": "1-2 hours",
                    "tat_stat": "30", "preparation": ""})
    with lab["app"].app_context():
        row = lab["db"].session.get(Investigation, test_id)
        assert (row.sample_type, row.tube, row.tat_min, row.tat_max) == ("دم", "أحمر", 60, 120)
        assert row.tat_stat_min == 30
    page = boss.get(f"/labs/tests/{test_id}/ranges").get_data(as_text=True)
    assert 'data-step="sample" data-done="1"' in page


def test_a_time_it_cannot_read_is_refused_not_guessed(lab):
    from app.models import Investigation

    with lab["app"].app_context():
        row = Investigation(name_ar="تحليل", kind="lab", tat_min=30, tat_max=30,
                            is_active=True)
        lab["db"].session.add(row)
        lab["db"].session.commit()
        test_id = row.id
    lab["sign_in"]("boss").post(f"/labs/tests/{test_id}/details",
                                data={"sample_type": "دم", "tat": "soon"})
    with lab["app"].app_context():
        row = lab["db"].session.get(Investigation, test_id)
        assert row.tat_min == 30 and row.sample_type is None
