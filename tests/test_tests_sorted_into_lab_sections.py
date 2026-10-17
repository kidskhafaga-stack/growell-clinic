"""كل تحليل في قسم من قايمة ثابتة — زي الدوا في عيلته.

* القسم غير «الفئة»: الفئة بتاعة الدكتور (كبد، التهابات) بتفضل زي ما هي؛
* البرنامج بيقترح لما يقدر يعرف — مزرعة = ميكروبيولوجي، تحليل بول = بول
  وبراز، فصيلة دم = بنك الدم، «تجلط» = تجلط — و«التهابات» ما بتقترحش حاجة
  (فيها CRP كيمياء وESR دم)؛ ومفيش حاجة بتتكتب غير بالحفظ؛
* القسم بيتختار من التعديل وبيتفلتر بيه، وبيتقري من الإكسيل وبيطلع فيه.
"""
import io
import os
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import pytest  # noqa: E402


@pytest.fixture()
def lab(clinic):
    from app.models import Investigation, Setting

    with clinic["app"].app_context():
        db = clinic["db"]
        Setting.set("mod_enabled:labs", "1")
        rows = {
            "urine": Investigation(name_ar="تحليل بول كامل", kind="lab", category="بول/كلى"),
            "culture": Investigation(name_ar="مزرعة بول", kind="lab", category="التهابات"),
            "group": Investigation(name_ar="فصيلة الدم", name_en="Blood group & Rh", kind="lab"),
            "pt": Investigation(name_ar="زمن البروثرومبين", kind="lab", category="تجلط"),
            "crp": Investigation(name_ar="بروتين سي التفاعلي", kind="lab", category="التهابات"),
        }
        db.session.add_all(rows.values())
        db.session.commit()
        clinic["ids"].update({k: v.id for k, v in rows.items()})
    return clinic


def _sections(lab):
    from app.models import Investigation

    with lab["app"].app_context():
        return {k: lab["db"].session.get(Investigation, i).section
                for k, i in lab["ids"].items() if k in ("urine", "culture", "group", "pt", "crp")}


def test_the_program_suggests_and_a_person_saves(lab):
    boss = lab["sign_in"]("boss")
    page = boss.get("/labs/tests/sections").get_data(as_text=True)
    ids = lab["ids"]
    for k, want in (("urine", "urine_stool"), ("culture", "microbiology"),
                    ("group", "blood_bank"), ("pt", "coagulation"), ("crp", "")):
        assert f'data-sort-row="{ids[k]}" data-suggested="{want}"' in page, k
    assert set(_sections(lab).values()) == {None}, "nothing until saved"
    form = {f"s_{ids[k]}": v for k, v in (("urine", "urine_stool"), ("culture", "microbiology"),
                                           ("group", "blood_bank"), ("pt", "coagulation"),
                                           ("crp", "chemistry"))}
    boss.post("/labs/tests/sections", data=form)
    got = _sections(lab)
    assert got == {"urine": "urine_stool", "culture": "microbiology", "group": "blood_bank",
                   "pt": "coagulation", "crp": "chemistry"}
    from app.models import Investigation
    with lab["app"].app_context():
        assert lab["db"].session.get(Investigation, ids["crp"]).category == "التهابات", \
            "the doctor's category is untouched"
    listing = boss.get("/labs/tests?section=chemistry").get_data(as_text=True)
    assert f'data-section-chip="chemistry"' in listing and "مزرعة بول" not in listing
    assert boss.get("/labs/tests").get_data(as_text=True).count("data-to-sections") == 1
    assert lab["sign_in"]("desk").get("/labs/tests/sections").status_code in (302, 403)


def test_a_sheet_brings_the_section_and_the_export_carries_it(lab):
    import openpyxl

    from app.models import Investigation
    from app.utils import lab_import

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Test Name", "Arabic Test Name", "Lab Section", "Component", "Unit"])
    sheet.append(["Vitamin D", "فيتامين د", "Vitamins & minerals", "25-OH Vitamin D", "ng/mL"])
    raw = io.BytesIO()
    book.save(raw)
    raw.seek(0)
    with lab["app"].app_context():
        plan = lab_import.read(raw)
        assert plan["tests"][0]["section"] == "vitamins"
        lab_import.apply(plan, {})
        lab["db"].session.commit()
        row = Investigation.query.filter_by(name_ar="فيتامين د").one()
        assert row.section == "vitamins"
        out = openpyxl.load_workbook(io.BytesIO(lab_import.export([row]))).active
        header = [c.value for c in out[1]]
        assert out.cell(row=2, column=header.index("Lab Section") + 1).value == "vitamins"
