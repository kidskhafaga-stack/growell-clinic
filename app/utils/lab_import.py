"""Bringing a laboratory's test list and its ranges in from a sheet.

Asked as *«هل الاستيراد هيبقى ذكي انه يفهم؟»* — and the answer is that it reads
whatever the laboratory sends, as long as a person can tell what a column is:

* **the columns are found by what they are called**, in Arabic or English and
  in any order — «Test Name» or «اسم التحليل», «Reference Low» or «الحد
  الأدنى»;
* **every sheet in the workbook is read** and the rows are put together by
  test: a list of names on one sheet, the ranges on another, the
  laboratory's tubes and times on a third — which is exactly how the first
  sheet this was built for arrived;
* **a test is recognised by any of its names**: «CBC» on the range sheet is
  the «Complete Blood Count (CBC)» of the list, because the list says so;
* **ages are read as written** — «15 days-4 weeks», «6-23 months», «18
  years+» — and kept as written beside the days they were read as.

**It reads; it never fills in.** A range the sheet does not have stays empty
and is said to be missing. The source a sheet gives for a figure is kept on
the figure. And whatever arrives is a *draft*: shown as a reference, never
used to call a result high or low until the laboratory's director approves
it (``approve_test``).

**Two steps, so nothing is written unseen.** ``read`` turns a workbook into a
plan and says what it found — what matches a test the program already has,
what is new, what looked wrong. ``apply`` writes that plan, with the links a
person confirmed, and nothing else. A test that already exists is linked and
**never renamed**: a doctor who orders «سكر صائم» keeps seeing «سكر صائم».
"""
import json
import os
import re
import uuid
from datetime import datetime, timedelta

from app.extensions import db

# --------------------------------------------------------------- headers ---
#: Every column the reader knows, and the names it answers to. Lower-cased and
#: trimmed before comparing.
HEADERS = {
    "test": ("test name", "test", "اسم التحليل", "التحليل"),
    "name_en": ("english test name", "english name", "الاسم الإنجليزي",
                "الاسم بالإنجليزي"),
    "name_ar": ("arabic test name", "arabic name", "الاسم العربي",
                "الاسم بالعربي", "اسم التحليل بالعربي"),
    "aliases": ("search aliases", "aliases", "أسماء البحث", "أسماء أخرى"),
    "category": ("category", "القسم", "الفئة"),
    "component": ("component", "analyte", "المكون", "المكوّن"),
    "unit": ("unit", "units", "الوحدة"),
    "age_from": ("age from", "من سن", "السن من"),
    "age_to": ("age to", "إلى سن", "الى سن", "السن إلى"),
    "sex": ("sex", "gender", "النوع", "الجنس"),
    "low": ("reference low", "low", "الحد الأدنى", "الحد الادنى"),
    "high": ("reference high", "high", "الحد الأعلى", "الحد الاعلى"),
    "ref_type": ("reference type", "نوع الحد"),
    "critical_low": ("critical low", "الحد الحرج الأدنى", "الحد الحرج الادنى"),
    "critical_high": ("critical high", "الحد الحرج الأعلى", "الحد الحرج الاعلى"),
    "specimen": ("specimen type", "specimen", "sample type", "نوع العينة"),
    "tube": ("tube color", "tube colour", "tube", "لون الأنبوبة", "الأنبوبة"),
    "tat": ("expected tat (routine)", "tat", "tat (routine)",
            "الوقت المتوقع", "الوقت المتوقع (عادي)"),
    "tat_stat": ("expected tat (stat)", "tat (stat)",
                 "الوقت المتوقع (مستعجل)"),
    "where": ("performed in / outsourced", "performed in", "in house / outsourced",
              "بيتعمل فين", "بيتعمل هنا ولا بره", "مكان التحليل"),
    "preparation": ("preparation", "التحضير"),
    "source": ("source", "المصدر"),
    "source_url": ("source url", "رابط المصدر"),
    "note": ("notes", "clinical / reference notes", "ملاحظات"),
}
_HEADER_OF = {alias: key for key, names in HEADERS.items() for alias in names}

#: A sheet is data when it names a test and at least one other known column.
_ENOUGH = 2


def _header_map(row):
    found = {}
    for i, cell in enumerate(row):
        key = _HEADER_OF.get(str(cell or "").strip().lower())
        if key and key not in found:
            found[key] = i
    names = {"test", "name_en", "name_ar"}
    # A name and a source is a page of notes about tests («handle neonatal
    # bilirubin by the hour»), not a list of them.
    facts = set(found) - names - {"source", "source_url", "note"}
    if not (names & set(found)) or len(found) < _ENOUGH or not facts:
        return None
    return found


# ----------------------------------------------------------------- values --
def _text(value):
    return str(value).strip() if value not in (None, "") else ""


def _number(value):
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


_DAYS = {"day": 1, "days": 1, "d": 1, "يوم": 1, "أيام": 1, "ايام": 1,
         "week": 7, "weeks": 7, "wk": 7, "wks": 7, "أسبوع": 7, "اسبوع": 7,
         "أسابيع": 7, "اسابيع": 7,
         "month": 30.4375, "months": 30.4375, "mo": 30.4375, "شهر": 30.4375,
         "شهور": 30.4375, "أشهر": 30.4375,
         "year": 365.25, "years": 365.25, "yr": 365.25, "yrs": 365.25,
         "y": 365.25, "سنة": 365.25, "سنه": 365.25, "سنين": 365.25,
         "سنوات": 365.25}
_WHOLE_LIFE = ("all ages", "all", "any", "كل الأعمار", "الكل")


def age_band(text):
    """``"6-23 months"`` → ``(from_days, to_days)``, the end exclusive and
    ``None`` for «and older». ``None`` when the band cannot be read.

    Read the way a laboratory writes it: both ends are whole units, so
    «6-23 months» runs to the day before 24 months, and «<1 year» or
    «1-<10 years» stop exactly at the figure.
    """
    s = re.sub(r"\s+", " ", (text or "").strip().lower())
    if not s:
        return None
    if s in _WHOLE_LIFE:
        return (0, None)
    if s in ("adult", "adults", "بالغ", "بالغين"):
        return (round(18 * 365.25), None)
    unit = r"([a-z؀-ۿ]+)"
    num = r"(\d+(?:\.\d+)?)"
    m = re.fullmatch(num + r" ?" + unit + r" ?\+", s) or \
        re.fullmatch(r"(?:>=|≥) ?" + num + r" ?" + unit, s)
    if m and m.group(2) in _DAYS:
        return (round(float(m.group(1)) * _DAYS[m.group(2)]), None)
    m = re.fullmatch(r"> ?" + num + r" ?" + unit, s)
    if m and m.group(2) in _DAYS:
        return (round((float(m.group(1)) + 1) * _DAYS[m.group(2)]), None)
    m = re.fullmatch(r"< ?" + num + r" ?" + unit, s)
    if m and m.group(2) in _DAYS:
        return (0, round(float(m.group(1)) * _DAYS[m.group(2)]))
    m = re.fullmatch(num + r" ?" + unit, s)
    if m and m.group(2) in _DAYS:
        per = _DAYS[m.group(2)]
        return (round(float(m.group(1)) * per),
                round((float(m.group(1)) + 1) * per))
    m = re.fullmatch(num + r" ?" + unit + r"? ?(?:-|–|to|إلى|الى) ?(<)? ?"
                     + num + r" ?" + unit, s)
    if m:
        first, first_unit, before, last, last_unit = m.groups()
        first_unit = first_unit or last_unit
        if first_unit in _DAYS and last_unit in _DAYS:
            start = round(float(first) * _DAYS[first_unit])
            end = float(last) if before else float(last) + 1
            return (start, round(end * _DAYS[last_unit]))
    return None


def minutes(text):
    """``"30-60"``, ``"2 hours"``, ``"1-2 days"`` → ``(min, max)`` minutes.

    A bare number is minutes, the way a laboratory agreement usually states
    it. ``None`` when it cannot be read — never a guess.
    """
    s = re.sub(r"\s+", " ", _text(text).lower())
    if not s:
        return None
    per = 1
    if re.search(r"hour|hr|\bh\b|ساع", s):
        per = 60
    elif re.search(r"day|يوم|أيام|ايام", s):
        per = 1440
    elif re.search(r"week|أسبوع|اسبوع", s):
        per = 10080
    found = [float(x) for x in re.findall(r"\d+(?:\.\d+)?", s)]
    if not found:
        return None
    low, high = found[0], found[1] if len(found) > 1 else found[0]
    return (round(low * per), round(high * per))


def _sex(text):
    s = _text(text).lower()
    if s in ("male", "m", "ذكر", "ذكور", "boy", "boys"):
        return "male"
    if s in ("female", "f", "أنثى", "انثى", "إناث", "girl", "girls"):
        return "female"
    return "all"


def _where(text):
    """``True`` done here, ``False`` sent out, ``None`` the sheet did not say."""
    s = _text(text).lower()
    if not s:
        return None
    if re.search(r"out|send|sent|external|referr|بره|برة|خارج", s):
        return False
    if re.search(r"in.?house|here|internal|هنا|داخل", s):
        return True
    return None


def _kind(ref_type, low, high):
    if low is None and high is None:
        return "note"
    if re.search(r"cutoff|cut-off|desirable|acceptable|minimum|target|إرشاد",
                 _text(ref_type).lower()):
        return "cutoff"
    return "interval"


# ------------------------------------------------------------ names/keys ---
_WORDS_LOOSE = {"count", "absolute", "serum", "level", "total", "test",
                "plasma", "blood"}
#: Words that name the paperwork, not the test — «Throat Swab Culture» is the
#: «Throat Culture».
_HARMLESS = {"swab", "tests", "test", "panel", "profile", "study", "screen"}


def key(text):
    """A name as the reader compares it: case, punctuation and spacing gone."""
    s = _text(text).lower()
    s = re.sub(r"[إأآا]", "ا", s).replace("ة", "ه").replace("ى", "ي")
    s = re.sub(r"[ً-ْ]", "", s)
    s = re.sub(r"[()\[\]/,.\-–:;*'’]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def loose(text):
    """A looser key, for telling that «Platelets» is the «Platelet Count»:
    plurals folded and the words that only say «a quantity of» dropped."""
    words = []
    for word in key(text).split():
        if word in _WORDS_LOOSE:
            continue
        if len(word) > 3 and word.endswith("s") and not word.endswith("ss"):
            word = word[:-1]
        words.append(word)
    return " ".join(words)


def _bracketed(text):
    m = re.search(r"\(([^)]+)\)", _text(text))
    return m.group(1) if m else ""


def _aliases(text):
    return [a.strip() for a in re.split(r"[;،,]", _text(text)) if a.strip()]


# ------------------------------------------------------------------ read ---
def read(stream):
    """Read a workbook into a plan (a plain dict, safe to keep as JSON)."""
    import openpyxl

    book = openpyxl.load_workbook(stream, read_only=True, data_only=True)
    tests = {}          # test key -> dict
    order = []
    index = {}          # any name -> test key
    range_rows = []     # (sheet, row number, values by header)
    problems = []
    sheets_read = []

    def remember(tkey, *names):
        for name in names:
            for k in (key(name), key(_bracketed(name))):
                if k:
                    index.setdefault(k, tkey)

    english = {}        # a test's own English name -> test key

    def primary(tkey, name_en):
        # The English name is the test's own and outranks any alias another
        # test happens to share.
        if key(name_en):
            index[key(name_en)] = tkey
            english.setdefault(key(name_en), tkey)

    for sheet in book.worksheets:
        header = None
        rows = sheet.iter_rows(values_only=True)
        for number, row in enumerate(rows, start=1):
            if header is None:
                header = _header_map(row)
                if header is None and number >= 3:
                    break           # not a data sheet: its top rows name nothing
                continue
            values = {k: row[i] if i < len(row) else None
                      for k, i in header.items()}
            if not any(_text(v) for v in values.values()):
                continue
            name_en = _text(values.get("name_en")) or _text(values.get("test"))
            name_ar = _text(values.get("name_ar"))
            if not (name_en or name_ar):
                continue
            if "name_en" in header or "name_ar" in header:
                # By the English name when there is one: two different tests
                # can share an «Arabic» name that is really the English one
                # (both 17-OHP tests), and must not be merged by it.
                # And by that name exactly, never through another test's
                # alias: a sheet that lists «Fecal Elastase» on its own and
                # as an alias of «Pancreatic Elastase» has listed two rows,
                # and merging them quietly is not the reader's call.
                tkey = (english.get(key(name_en)) if name_en
                        else index.get(key(name_ar)))
                if tkey is None:
                    tkey = key(name_en or name_ar)
                    tests[tkey] = _blank_test(name_en or name_ar)
                    order.append(tkey)
                test = tests[tkey]
                test["name_en"] = test["name_en"] or name_en
                test["name_ar"] = test["name_ar"] or name_ar
                for alias in _aliases(values.get("aliases")):
                    if alias not in test["aliases"]:
                        test["aliases"].append(alias)
                remember(tkey, name_en, name_ar, *test["aliases"])
                primary(tkey, name_en)
                _test_facts(test, values)
            has_range = any(_text(values.get(k)) for k in
                            ("low", "high", "age_from", "critical_low",
                             "critical_high"))
            if _text(values.get("component")) and (
                    has_range or "name_en" not in header):
                range_rows.append((sheet.title, number, values))
        if header is not None:
            sheets_read.append(sheet.title)

    analytes = {}       # analyte key -> dict
    exact_index = {}    # a name or alias, tidied -> analyte key
    loose_index = {}    # a loose name -> analyte keys
    # An alias two of the sheet's tests both claim («PT» on «Prothrombin
    # Time» and on «Prothrombin Time / INR») points at neither: guessing
    # would give PT the INR's ranges.
    claimed = {}
    for tkey in order:
        test = tests[tkey]
        for alias in test["aliases"]:
            claimed.setdefault(key(alias), set()).add(tkey)
    shared = {k for k, owners in claimed.items() if len(owners) > 1}

    def analyte_for(name, unit=None, aliases=(), part=False):
        """The analyte this name is, made if there is none.

        A test of the sheet's is matched **only** by its exact name or an
        alias the sheet gives — «Reticulocyte Count» (a percentage) and
        «Absolute Reticulocyte Count» (a number) are two analytes, and a
        looser rule would have given one the other's ranges. A panel's part
        («Platelets» of a CBC) may also be matched loosely, and only when the
        loose name points at one analyte and no other.
        """
        for k in [key(name)] + [key(a) for a in aliases if key(a) not in shared]:
            if k and k in exact_index:
                found = analytes[exact_index[k]]
                found["unit"] = found["unit"] or _text(unit)
                return exact_index[k]
        if part:
            near = loose_index.get(loose(name), set())
            if len(near) == 1:
                akey = next(iter(near))
                analytes[akey]["unit"] = analytes[akey]["unit"] or _text(unit)
                exact_index.setdefault(key(name), akey)
                return akey
        akey = key(name)
        while akey in analytes:
            akey += "'"
        analytes[akey] = {"name": _text(name), "unit": _text(unit),
                          "aliases": list(aliases), "ranges": [],
                          "critical_low": None, "critical_high": None}
        exact_index[key(name)] = akey
        for k in [key(a) for a in aliases if key(a) not in shared]:
            if k:
                exact_index.setdefault(k, akey)
        loose_index.setdefault(loose(name), set()).add(akey)
        return akey

    # The sheet's own tests first, so a panel's «Hemoglobin» finds the
    # «Hemoglobin» test's analyte rather than making a second one.
    own = {}
    for tkey in order:
        test = tests[tkey]
        own[tkey] = analyte_for(test["name_en"] or test["name_ar"] or test["label"],
                                aliases=test["aliases"])

    seen_ranges = {}
    for sheet, number, values in range_rows:
        tname = _text(values.get("test")) or _text(values.get("name_en")) \
            or _text(values.get("name_ar"))
        tkey = index.get(key(tname)) or index.get(key(_bracketed(tname))) \
            or index.get(loose(tname))
        if tkey is None:
            tkey = key(tname)
            tests[tkey] = _blank_test(tname)
            order.append(tkey)
            remember(tkey, tname)
        test = tests[tkey]
        _test_facts(test, values)
        component = _text(values.get("component"))
        if key(component) in _PLACEHOLDERS:
            # «Panel / Profile»: the sheet says the test has parts and does
            # not list them. Said in the preview, never guessed.
            test["parts_missing"] = True
            continue
        low, high = _number(values.get("low")), _number(values.get("high"))
        crit_low = _number(values.get("critical_low"))
        crit_high = _number(values.get("critical_high"))
        age_text = _text(values.get("age_from"))
        whole = key(component) in {key(tname), key(test["name_en"]),
                                   key(test["name_ar"]), key(test["label"])}
        if whole and not age_text and low is None and high is None \
                and crit_low is None and crit_high is None:
            continue                # «the whole test», not one of its parts
        akey = analyte_for(component, values.get("unit"), part=True)
        if akey not in test["analytes"]:
            test["analytes"].append(akey)
        analyte = analytes[akey]
        if _text(values.get("age_to")):
            age_text = f"{age_text}-{_text(values.get('age_to'))}"
        if not age_text and low is None and high is None:
            # A row that only carries the laboratory's alert limits for the
            # analyte as a whole.
            analyte["critical_low"] = crit_low if crit_low is not None \
                else analyte["critical_low"]
            analyte["critical_high"] = crit_high if crit_high is not None \
                else analyte["critical_high"]
            continue
        band = age_band(age_text) if age_text else (0, None)
        where = f"{sheet} · {number}"
        if band is None:
            problems.append({"where": where, "what": "age",
                             "detail": f"{tname} · {component} · «{age_text}»"})
            band = (0, None)
        sex = _sex(values.get("sex"))
        kind = _kind(values.get("ref_type"), low, high)
        if low is not None and high is not None and low > high:
            problems.append({"where": where, "what": "low_above_high",
                             "detail": f"{tname} · {component} · {low} > {high}"})
            continue
        fingerprint = (akey, band, sex, low, high, kind)
        if fingerprint in seen_ranges:
            # The same row on two sheets: one may say where it came from and
            # the other not, and the one that says wins.
            earlier = seen_ranges[fingerprint]
            for field, source in (("source", "source"),
                                  ("source_url", "source_url"),
                                  ("note", "note")):
                if not earlier[field] and _text(values.get(source)):
                    earlier[field] = _text(values.get(source))[:300]
            continue
        entry = {
            "from": band[0], "to": band[1], "label": age_text or "",
            "sex": sex, "kind": kind, "low": low, "high": high,
            "critical_low": crit_low, "critical_high": crit_high,
            "note": _text(values.get("note"))[:255] or None,
            "source": _text(values.get("source"))[:160] or None,
            "source_url": _text(values.get("source_url"))[:300] or None,
        }
        seen_ranges[fingerprint] = entry
        analyte["ranges"].append(entry)

    for tkey in order:
        test = tests[tkey]
        if not test["analytes"]:
            test["analytes"].append(
                own.get(tkey) or analyte_for(test["label"]))
    # Only what some test measures — a panel's own name is not an analyte.
    used = {a for t in tests.values() for a in t["analytes"]}
    analytes = {k: v for k, v in analytes.items() if k in used}
    problems.extend(_gaps(analytes))
    return {"tests": [dict(tests[k], key=k) for k in order],
            "analytes": analytes, "problems": problems,
            "sheets": sheets_read}


#: What a sheet writes in the «component» column when it means «this test
#: has parts» without naming them.
_PLACEHOLDERS = {"panel profile", "panel", "profile", "panel profil", "n a",
                 "na", "-", "multiple", "several"}


def _blank_test(name):
    return {"name_en": "", "name_ar": "", "aliases": [], "category": "",
            "parts_missing": False,
            "specimen": "", "tube": "", "tat": None, "tat_stat": None,
            "in_house": None, "preparation": "", "analytes": [],
            "label": _text(name)}


def _test_facts(test, values):
    """The test-level facts a row carries, first one given wins."""
    for field, source in (("category", "category"), ("specimen", "specimen"),
                          ("tube", "tube"), ("preparation", "preparation")):
        if not test[field] and _text(values.get(source)):
            test[field] = _text(values.get(source))
    for field, source in (("tat", "tat"), ("tat_stat", "tat_stat")):
        if test[field] is None and _text(values.get(source)):
            test[field] = minutes(values.get(source))
    if test["in_house"] is None:
        test["in_house"] = _where(values.get("where"))


def _gaps(analytes):
    """Bands of one analyte and sex that leave an age uncovered between two
    rows — said, never filled."""
    out = []
    for akey, analyte in analytes.items():
        by_sex = {}
        for row in analyte["ranges"]:
            if row["kind"] == "note":
                continue
            by_sex.setdefault(row["sex"], []).append(row)
        for sex, rows in by_sex.items():
            rows = sorted(rows, key=lambda r: r["from"])
            for one, two in zip(rows, rows[1:]):
                if one["to"] is not None and two["from"] > one["to"] + 1:
                    out.append({"where": analyte["name"], "what": "age_gap",
                                "detail": f"{analyte['name']} · {sex} · "
                                          f"{one['label']} → {two['label']}"})
    return out


# --------------------------------------------------------------- matching --
def match(plan):
    """How the sheet's tests meet the catalogue the program already has.

    Returns ``{"auto": {test key: id}, "questions": [...]}``.

    * **auto** — a sheet test whose name or alias is, once tidied, the name,
      alias or code of exactly one of ours, and no other sheet test claims
      the same one. Linked without asking: «CBC» is «CBC».
    * **questions** — one per test of ours that found no such twin, with the
      sheet's closest names offered and the closest preselected when one is
      clearly closest. «سكر صائم / Fasting Blood Sugar» meets «Fasting Blood
      Glucose» here, and a person says whether it is the same test. Asked
      from our side because ours is the short list: a sheet of five hundred
      would otherwise ask five hundred questions.
    """
    from app.models import Investigation

    # Rows with a code first: a specialty panel names a test by its code,
    # so where a clinic has the same test twice, that is the one to keep.
    catalogue = (Investigation.query.filter(Investigation.kind == "lab")
                 .order_by(Investigation.code.is_(None), Investigation.id)
                 .all())
    ours_by_name, ours_by_en = {}, {}
    for row in catalogue:
        for name in [row.name_ar, row.name_en, row.code, _bracketed(row.name_ar),
                     _bracketed(row.name_en)] + _aliases(row.aliases):
            if key(name):
                ours_by_name.setdefault(key(name), set()).add(row.id)
        if key(row.name_en):
            ours_by_en.setdefault(key(row.name_en), set()).add(row.id)
    claims = {}
    for test in plan["tests"]:
        # The English name first: it is the test's own. Only when it meets
        # nothing do the Arabic name and the aliases get a say — two tests
        # can share an «Arabic» name that is really the English one.
        found = set(ours_by_en.get(key(test["name_en"]), set()))
        if not found:
            for name in [test["name_en"], test["name_ar"], test["label"],
                         _bracketed(test["name_en"])] + test["aliases"]:
                found |= ours_by_name.get(key(name), set())
        if len(found) == 1:
            claims.setdefault(next(iter(found)), []).append(test["key"])
    auto = {keys[0]: cid for cid, keys in claims.items() if len(keys) == 1}
    linked = set(auto.values())

    tokens = {t["key"]: set(loose(t["name_en"] or t["label"]).split()) - _HARMLESS
              for t in plan["tests"]}
    label = {t["key"]: (t["name_en"] or t["label"]) for t in plan["tests"]}
    twins = {}
    for test in plan["tests"]:
        for name in [test["name_en"], test["name_ar"], test["label"],
                     _bracketed(test["name_en"])] + test["aliases"]:
            for cid in ours_by_name.get(key(name), set()):
                twins.setdefault(cid, [])
                if test["key"] not in twins[cid]:
                    twins[cid].append(test["key"])
    taken = set(auto)
    picked = set()
    questions = []
    for row in catalogue:
        if row.id in linked:
            continue
        mine = (set(loose(row.name_en).split()) |
                set(loose(_bracketed(row.name_ar)).split())) - _HARMLESS
        # Same name, alias or code, claimed by more than one side: the
        # likeliest answer, offered first — including a sheet test another
        # of ours already took, because a clinic that has CRP twice has one
        # CRP, and the second copy is that same test.
        same = list(twins.get(row.id, []))
        scored = []
        for tkey, theirs in tokens.items():
            if tkey in taken or tkey in same or not mine or not theirs:
                continue
            common = len(mine & theirs)
            if common:
                scored.append((common / len(mine | theirs), tkey,
                               mine <= theirs))
        scored.sort(key=lambda x: (-x[0], label[x[1]]))
        best = same + [k for _, k, _ in scored][:max(0, 4 - len(same))]
        # Preselected only when it is plainly the same test: a twin by name,
        # or the one sheet name that holds every word of ours that matters —
        # «Serum Calcium» → «Total Calcium», never «Sweat Chloride» →
        # «Chloride».
        pick = None
        already = [k for k in same if k in taken]
        if len(same) > 1 and len(already) == 1:
            # «CRP» meets «C-Reactive Protein» and «Neonatal C-Reactive
            # Protein»; the first is the one our other CRP was linked to.
            same = already + [k for k in same if k not in already]
            pick = already[0]
            questions.append({"id": row.id, "name_ar": row.name_ar,
                              "name_en": row.name_en or "",
                              "candidates": [[k, label[k]] for k in same[:4]],
                              "pick": pick})
            continue
        if len(same) == 1:
            # A twin by name — also for the clinic's second copy of the same
            # test («صورة دم كاملة» and «صورة دم كاملة (CBC)»): both are that
            # test and both get what it measures.
            pick = same[0]
            questions.append({"id": row.id, "name_ar": row.name_ar,
                              "name_en": row.name_en or "",
                              "candidates": [[k, label[k]] for k in best],
                              "pick": pick})
            continue
        elif not same:
            whole = [k for score, k, covers in scored if covers]
            if len(whole) == 1 or (whole and len(scored) > 1
                                   and scored[0][1] == whole[0]
                                   and scored[0][0] > scored[1][0]):
                pick = whole[0]
        if pick in picked:
            pick = None
        if pick:
            picked.add(pick)
        questions.append({"id": row.id, "name_ar": row.name_ar,
                          "name_en": row.name_en or "",
                          "candidates": [[k, label[k]] for k in best],
                          "pick": pick})
    return {"auto": auto, "questions": questions}


def links_from(plan, answers):
    """The final ``{test key: [catalogue ids]}``: the automatic links, then
    each question as answered (``answers`` is ``{catalogue id: test key or
    ""}``). A catalogue row is linked to one sheet test at most; a sheet test
    may be linked to several rows of ours when the clinic has it twice."""
    found = match(plan)
    links = {k: [cid] for k, cid in found["auto"].items()}
    used = set(found["auto"].values())
    keys = {t["key"] for t in plan["tests"]}
    for q in found["questions"]:
        tkey = answers.get(str(q["id"]), answers.get(q["id"]))
        if tkey and tkey in keys and q["id"] not in used:
            links.setdefault(tkey, []).append(q["id"])
            used.add(q["id"])
    return links


# ------------------------------------------------------------- keep/load ---
def _folder():
    from flask import current_app

    path = os.path.join(current_app.instance_path, "lab_imports")
    os.makedirs(path, exist_ok=True)
    return path


def keep(plan):
    """Put a read plan aside until somebody presses «import». Returns its id.
    Plans older than a day are cleared on the way."""
    folder = _folder()
    cutoff = datetime.utcnow() - timedelta(days=1)
    for name in os.listdir(folder):
        full = os.path.join(folder, name)
        if datetime.utcfromtimestamp(os.path.getmtime(full)) < cutoff:
            os.remove(full)
    token = uuid.uuid4().hex
    with open(os.path.join(folder, f"{token}.json"), "w", encoding="utf-8") as fh:
        json.dump(plan, fh, ensure_ascii=False)
    return token


def load(token):
    if not re.fullmatch(r"[0-9a-f]{32}", token or ""):
        return None
    path = os.path.join(_folder(), f"{token}.json")
    if not os.path.exists(path):
        return None
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def forget(token):
    path = os.path.join(_folder(), f"{token}.json")
    if re.fullmatch(r"[0-9a-f]{32}", token or "") and os.path.exists(path):
        os.remove(path)


# ------------------------------------------------------------------ apply ---
def apply(plan, links, user=None, show_new=True):
    """Write a plan. ``links`` is ``{test key: [catalogue ids]}`` — the
    automatic links and the questions a person answered (``links_from``);
    every other test in the plan is new. Returns counts.

    What it never does: rename a test that exists, change a test's price,
    touch an order or a result, or overwrite a range the laboratory has
    approved — incoming figures for such an analyte wait as drafts beside it.
    """
    from app.models import (Investigation, LabAnalyte, LabRange,
                            LabTestAnalyte)

    counts = {"linked": 0, "created": 0, "analytes": 0, "ranges": 0}
    analyte_ids = {}
    by_name, by_alias = {}, {}
    for row in LabAnalyte.query.all():
        by_name.setdefault(key(row.name), row.id)
        for alias in _aliases(row.aliases):
            by_alias.setdefault(key(alias), set()).add(row.id)
    cleared = set()
    for akey, incoming in plan["analytes"].items():
        # The same analyte by its own name first; by an alias only when the
        # alias belongs to one analyte and no other («INR» must not find
        # «Prothrombin Time / INR»).
        aid = by_name.get(key(incoming["name"]))
        if aid is None:
            for alias in [incoming["name"]] + incoming["aliases"]:
                owners = by_alias.get(key(alias), set())
                if len(owners) == 1:
                    aid = next(iter(owners))
                    break
        if aid is None:
            row = LabAnalyte(name=incoming["name"][:160],
                             unit=(incoming["unit"] or None),
                             aliases="; ".join(incoming["aliases"])[:400] or None)
            db.session.add(row)
            db.session.flush()
            aid = row.id
            by_name.setdefault(key(row.name), aid)
            for alias in incoming["aliases"]:
                by_alias.setdefault(key(alias), set()).add(aid)
            counts["analytes"] += 1
        else:
            row = db.session.get(LabAnalyte, aid)
            row.unit = row.unit or incoming["unit"] or None
        analyte_ids[akey] = aid
        if incoming["ranges"]:
            # The drafts an earlier sheet left are replaced — once per
            # analyte, so two of this sheet's names for one analyte add up
            # rather than the second wiping the first.
            if aid not in cleared:
                (LabRange.query.filter(LabRange.analyte_id == aid,
                                       LabRange.approved_at.is_(None))
                 .delete(synchronize_session=False))
                cleared.add(aid)
            for r in incoming["ranges"]:
                db.session.add(LabRange(
                    analyte_id=aid, age_from_days=r["from"],
                    age_to_days=r["to"], sex=r["sex"], kind=r["kind"],
                    low=r["low"], high=r["high"],
                    critical_low=(r["critical_low"] if r["critical_low"] is not None
                                  else incoming["critical_low"]),
                    critical_high=(r["critical_high"] if r["critical_high"] is not None
                                   else incoming["critical_high"]),
                    age_label=(r["label"] or "")[:60] or None, note=r["note"],
                    source=r["source"], source_url=r["source_url"]))
                counts["ranges"] += 1

    for test in plan["tests"]:
        targets = [row for row in (db.session.get(Investigation, int(cid))
                                   for cid in links.get(test["key"], []))
                   if row is not None]
        if targets:
            counts["linked"] += len(targets)
        else:
            target = Investigation(
                name_ar=(test["name_ar"] or test["name_en"] or test["label"])[:160],
                # A panel only the ranges sheet names («Coagulation») has
                # no list row; its name there is its English name.
                name_en=(test["name_en"] or test["label"] or "")[:160] or None,
                kind="lab", category=(test["category"] or None) and
                test["category"][:80], is_active=bool(show_new))
            db.session.add(target)
            db.session.flush()
            targets = [target]
            counts["created"] += 1
        for target in targets:
            _fill_blanks(target, test)
            have = {link.analyte_id for link in target.analyte_links}
            place = len(have)
            for akey in test["analytes"]:
                aid = analyte_ids[akey]
                if aid in have:
                    continue
                db.session.add(LabTestAnalyte(investigation_id=target.id,
                                              analyte_id=aid, sort_order=place))
                have.add(aid)
                place += 1
            db.session.flush()
    return counts


def _fill_blanks(target, test):
    """Only ever fill what is empty — the name a clinic chose, its unit, its
    price and its category are its own."""
    known = set(_aliases(target.aliases))
    extra = [a for a in test["aliases"] if a not in known
             and key(a) not in (key(target.name_ar), key(target.name_en))]
    if extra:
        target.aliases = "; ".join(list(known) + extra)[:400]
    if not target.name_en and test["name_en"]:
        target.name_en = test["name_en"][:160]
    if not target.sample_type and test["specimen"]:
        target.sample_type = test["specimen"][:40]
    if not target.tube and test["tube"]:
        target.tube = test["tube"][:60]
    if not target.preparation and test["preparation"]:
        target.preparation = test["preparation"][:255]
    if target.tat_min is None and test["tat"]:
        target.tat_min, target.tat_max = test["tat"]
    if target.tat_stat_min is None and test["tat_stat"]:
        target.tat_stat_min, target.tat_stat_max = test["tat_stat"]
    # Where it is done is the laboratory's to say, and said only when the
    # sheet says it: an empty cell leaves the clinic's own answer alone.
    if test["in_house"] is not None:
        target.in_house = test["in_house"]


# ---------------------------------------------------------------- approve ---
def approve_test(investigation, user):
    """The laboratory's director approves the ranges of one test's analytes.

    Where an analyte has drafts, they replace whatever was approved before —
    that is what a new sheet from the laboratory means. Returns how many rows
    are now approved.
    """
    from app.models import LabRange

    now = datetime.utcnow()
    done = 0
    for link in investigation.analyte_links:
        drafts = [r for r in link.analyte.ranges if not r.approved]
        if not drafts:
            continue
        for old in [r for r in link.analyte.ranges if r.approved]:
            db.session.delete(old)
        for row in drafts:
            row.approved_at = now
            row.approved_by = getattr(user, "id", None)
            done += 1
    db.session.flush()
    return done


# ----------------------------------------------------------------- export ---
EXPORT_COLUMNS = ("Test Name", "Arabic Test Name", "Search Aliases", "Category",
                  "Component", "Unit", "Age From", "Sex", "Reference Low",
                  "Reference High", "Reference Type", "Critical Low",
                  "Critical High", "Specimen Type", "Tube Color",
                  "Expected TAT (Routine)", "Expected TAT (STAT)",
                  "Performed In / Outsourced", "Preparation", "Source",
                  "Source URL", "Status")


def export(rows):
    """The laboratory's catalogue as a sheet it can complete and send back —
    the same columns ``read`` understands. One line per range, and one per
    analyte that has none yet, so an empty cell is a question put to the
    laboratory, not a gap nobody sees."""
    import io

    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Laboratory_Tests"
    sheet.append(EXPORT_COLUMNS)

    def span(low, high):
        if low is None:
            return ""
        return f"{low}-{high}" if high not in (None, low) else f"{low}"

    for test in rows:
        facts = [test.name_en or test.name_ar, test.name_ar, test.aliases or "",
                 test.category or ""]
        where = ("" if test.in_house is None else
                 ("In-house" if test.in_house else "Outsourced"))
        tail = [test.sample_type or "", test.tube or "",
                span(test.tat_min, test.tat_max),
                span(test.tat_stat_min, test.tat_stat_max), where,
                test.preparation or ""]
        links = list(test.analyte_links) or [None]
        for link in links:
            analyte = link.analyte if link else None
            ranges = analyte.ranges if analyte else []
            if not ranges:
                sheet.append(facts + [analyte.name if analyte else
                                      (test.name_en or test.name_ar),
                                      (analyte.unit if analyte else test.unit) or "",
                                      "", "", "", "", "", "", ""] + tail +
                             ["", "", ""])
                continue
            for r in ranges:
                sheet.append(facts + [
                    analyte.name, analyte.unit or "", r.age_label or "",
                    r.sex.capitalize() if r.sex != "all" else "All",
                    r.low if r.low is not None else "",
                    r.high if r.high is not None else "", r.kind,
                    r.critical_low if r.critical_low is not None else "",
                    r.critical_high if r.critical_high is not None else ""]
                    + tail + [r.source or "", r.source_url or "",
                              "approved" if r.approved else "draft"])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
